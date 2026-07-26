"""
import_hot_questions.py

Читає .docx файли "Гарячих питань" (за шаблоном ШАБЛОН ГАРЯЧЕ ПИТАННЯ.docx)
з вказаної папки і імпортує їх у базу даних (таблиця questions).

Використання:
    python import_hot_questions.py /шлях/до/папки/з/файлами

Для кожного файлу скрипт:
  1. Розпізнає назву і 7 блоків контенту за emoji-маркерами.
  2. Показує назву і просить вибрати тему зі списку (одна цифра).
  3. Генерує унікальний slug з назви.
  4. Записує все в базу даних.

Наприкінці виводить підсумок: скільки файлів імпортовано, скільки пропущено.
"""

import sys
import os
import re
import unicodedata

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

import database  # використовує get_connection, USE_POSTGRES з вашого database.py


# ──────────────────────────────────────────────────────────────────────────
# Маркери блоків шаблону (мають збігатись з реальними emoji у ваших файлах)
# ──────────────────────────────────────────────────────────────────────────

MARKERS = {
    "👤": "block_audience",
    "⚠️": "block_problem",
    "✅": "block_solution",
    "📋": "block_example",
    "❌": "block_mistakes",
    "📌": "block_checklist",
    "📎": "block_sources",
}
TITLE_MARKER = "🔥"


# ──────────────────────────────────────────────────────────────────────────
# Допоміжне: обхід документа в порядку "як написано" (параграфи + таблиці)
# ──────────────────────────────────────────────────────────────────────────

def iter_block_items(document):
    """Повертає параграфи і таблиці документа у природному порядку читання."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def table_text(table):
    """Витягує весь текст з таблиці, рядок за рядком, через новий рядок."""
    lines = []
    for row in table.rows:
        for cell in row.cells:
            text = cell.text.strip()
            if text:
                lines.append(text)
    return "\n".join(lines)


# ──────────────────────────────────────────────────────────────────────────
# Парсинг одного .docx файлу
# ──────────────────────────────────────────────────────────────────────────

def parse_docx(filepath):
    """
    Повертає словник {title, block_audience, block_problem, ...}
    або None, якщо файл не вдалось розпізнати за шаблоном.
    """
    doc = Document(filepath)
    items = list(iter_block_items(doc))

    title = None
    blocks = {key: "" for key in MARKERS.values()}
    current_key = None

    for item in items:
        if isinstance(item, Table):
            text = table_text(item)
            if not text:
                continue
            if title is None and TITLE_MARKER in text:
                # Перша таблиця з 🔥 — це заголовок. Прибираємо маркер і зайві слова.
                title = text.replace(TITLE_MARKER, "").strip()
                title = re.sub(r"^ГАРЯЧЕ ПИТАННЯ\s*\d*\s*", "", title, flags=re.IGNORECASE).strip()
                continue
            if current_key == "block_example":
                blocks["block_example"] += ("\n" if blocks["block_example"] else "") + text
                continue
            # Будь-яка інша таблиця без відомого контексту — додаємо до поточного блоку
            if current_key:
                blocks[current_key] += ("\n" if blocks[current_key] else "") + text

        elif isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue

            matched_marker = None
            for marker, key in MARKERS.items():
                if text.startswith(marker):
                    matched_marker = key
                    break

            if matched_marker:
                current_key = matched_marker
                # Прибираємо номер блоку типу "1. Кому це актуально" з маркера
                cleaned = text
                for marker in MARKERS:
                    cleaned = cleaned.replace(marker, "")
                cleaned = re.sub(r"^\s*\d+\.\s*", "", cleaned).strip()
                # Якщо після заголовка блоку в тому ж рядку є вже якийсь текст — додати
                if cleaned and not cleaned.lower().startswith((
                    "кому це актуально", "суть проблеми", "як діяти правильно",
                    "приклад з практики", "типові помилки", "короткий чеклист",
                    "джерело", "джерело / нормативна база"
                )):
                    blocks[current_key] += ("\n" if blocks[current_key] else "") + cleaned
                continue

            if current_key:
                blocks[current_key] += ("\n" if blocks[current_key] else "") + text

    if not title or not any(blocks.values()):
        return None

    result = {"title": title}
    result.update(blocks)
    return result


# ──────────────────────────────────────────────────────────────────────────
# Генерація slug (транслітерація укр -> латиниця, без спецсимволів)
# ──────────────────────────────────────────────────────────────────────────

TRANSLIT_MAP = {
    "а": "a", "б": "b", "в": "v", "г": "h", "ґ": "g", "д": "d", "е": "e",
    "є": "ie", "ж": "zh", "з": "z", "и": "y", "і": "i", "ї": "i", "й": "i",
    "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r",
    "с": "s", "т": "t", "у": "u", "ф": "f", "х": "kh", "ц": "ts", "ч": "ch",
    "ш": "sh", "щ": "shch", "ь": "", "ю": "iu", "я": "ia", "'": "", "’": "",
}


def slugify(text, max_words=5):
    text = text.lower()
    out = []
    for ch in text:
        out.append(TRANSLIT_MAP.get(ch, ch))
    text = "".join(out)
    text = re.sub(r"[^a-z0-9\s-]", "", text)
    words = text.split()[:max_words]
    slug = "-".join(words)
    slug = re.sub(r"-+", "-", slug).strip("-")
    return slug or "pytannya"


def make_unique_slug(base_slug, existing_slugs):
    slug = base_slug
    counter = 2
    while slug in existing_slugs:
        slug = f"{base_slug}-{counter}"
        counter += 1
    existing_slugs.add(slug)
    return slug


# ──────────────────────────────────────────────────────────────────────────
# Робота з базою: список тем, вставка питання
# ──────────────────────────────────────────────────────────────────────────

def load_topics():
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("SELECT id, name FROM topics WHERE is_active = 1 ORDER BY sort_order, name")
    rows = cur.fetchall()
    conn.close()
    # Уніфікуємо доступ (sqlite3.Row і psycopg2 RealDictRow обидва підтримують ['name'])
    return [(row["id"], row["name"]) for row in rows]


def get_existing_slugs():
    conn = database.get_connection()
    cur = conn.cursor()
    cur.execute("SELECT slug FROM questions")
    rows = cur.fetchall()
    conn.close()
    return {row["slug"] for row in rows}


def insert_question(topic_id, slug, parsed, price=99):
    conn = database.get_connection()
    cur = conn.cursor()
    placeholder = "%s" if database.USE_POSTGRES else "?"
    columns = [
        "topic_id", "slug", "title",
        "block_audience", "block_problem", "block_solution",
        "block_example", "block_mistakes", "block_checklist", "block_sources",
        "price",
    ]
    values = [
        topic_id, slug, parsed["title"],
        parsed["block_audience"], parsed["block_problem"], parsed["block_solution"],
        parsed["block_example"], parsed["block_mistakes"], parsed["block_checklist"],
        parsed["block_sources"], price,
    ]
    placeholders = ", ".join([placeholder] * len(columns))
    sql = f"INSERT INTO questions ({', '.join(columns)}) VALUES ({placeholders})"
    cur.execute(sql, values)
    conn.commit()
    conn.close()


# ──────────────────────────────────────────────────────────────────────────
# Основний сценарій
# ──────────────────────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print("Використання: python import_hot_questions.py /шлях/до/папки")
        sys.exit(1)

    folder = sys.argv[1]
    if not os.path.isdir(folder):
        print(f"Папку не знайдено: {folder}")
        sys.exit(1)

    docx_files = [f for f in os.listdir(folder) if f.lower().endswith(".docx") and not f.startswith("~$")]
    if not docx_files:
        print("У цій папці не знайдено жодного .docx файлу.")
        sys.exit(1)

    topics = load_topics()
    if not topics:
        print("У базі немає жодної теми. Спершу запустіть database.py.")
        sys.exit(1)

    print("\nДоступні теми:")
    for idx, (topic_id, name) in enumerate(topics, start=1):
        print(f"  {idx}. {name}")

    existing_slugs = get_existing_slugs()

    imported = 0
    skipped = 0

    for filename in sorted(docx_files):
        filepath = os.path.join(folder, filename)
        print(f"\n{'-'*60}")
        print(f"Файл: {filename}")

        parsed = parse_docx(filepath)
        if parsed is None:
            print("  ⚠️  Не вдалось розпізнати структуру файлу — пропускаю.")
            skipped += 1
            continue

        print(f"  Назва: {parsed['title']}")

        while True:
            choice = input(f"  Оберіть тему (1-{len(topics)}), або 's' щоб пропустити файл: ").strip().lower()
            if choice == "s":
                print("  Пропущено.")
                skipped += 1
                break
            if choice.isdigit() and 1 <= int(choice) <= len(topics):
                topic_id, topic_name = topics[int(choice) - 1]
                slug = make_unique_slug(slugify(parsed["title"]), existing_slugs)
                insert_question(topic_id, slug, parsed)
                print(f"  ✅ Імпортовано в тему «{topic_name}» (slug: {slug})")
                imported += 1
                break
            print("  Некоректний вибір, спробуйте ще раз.")

    print(f"\n{'='*60}")
    print(f"Готово. Імпортовано: {imported}. Пропущено: {skipped}.")


if __name__ == "__main__":
    main()
