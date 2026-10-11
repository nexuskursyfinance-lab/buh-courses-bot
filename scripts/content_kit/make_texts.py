# -*- coding: utf-8 -*-
"""Тексти публікацій (docx) і пакет Viber/WhatsApp (md + картинки) для партії ГП.
Запуск: python3 make_texts.py <папка_виводу> <назва_партії, напр. 21-30> <дата>"""
import os, sys, shutil
from docx import Document
from docx.shared import Pt, RGBColor
from content import BRANCH as BR, GP

GOLD = RGBColor(*BR["color"])


def link(n):
    return f"https://t.me/{BR['bot']}?start=gp{n:02d}"


def tag_title(q):
    a, b = q["tag"].split(" · ", 1)
    return f"{a} · {b.lower()}"


def docx_texts(out, batch, date):
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(11)
    p = d.add_paragraph()
    r = p.add_run(f"{BR['name'].capitalize()} · тексти публікацій ГП-{batch} · {date}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(14), GOLD
    for line in ["Посилання вже вставлено (deep link на ГП у боті).",
                 f"Картки й відео — 7_ПОСТИ/ГП {batch.replace('-', '-')}/1_ГП-{batch.replace('-', '–')} картки і відео.",
                 "Instagram: замість посилання — «посилання в шапці профілю».",
                 "Публікуємо лише після явного «так» на кожен пост."]:
        d.add_paragraph(line)
    for q in GP:
        code = f"ГП-{q['n']:02d}"
        h = d.add_heading(f"{code} · {tag_title(q)}", level=2)
        for r in h.runs:
            r.font.color.rgb = GOLD
        d.add_paragraph(f"Файли: {code}_картка-1.png, {code}_картка-2.png, {code}_картка-3.png, {code}_відео_9x16.mp4")
        d.add_paragraph().add_run("Пост — Telegram / Facebook / Instagram").bold = True
        for para in q["post"].split("\n"):
            d.add_paragraph(para)
        d.add_paragraph(link(q["n"]))
        d.add_paragraph(q["tags"])
        d.add_paragraph().add_run("Підпис — Reels / TikTok / Shorts").bold = True
        d.add_paragraph(q["reel"])
        d.add_paragraph("Розбір — у боті за посиланням у профілі")
        d.add_paragraph(q["tags"])
    path = os.path.join(out, f"Тексти публікацій_ГП-{batch}.docx")
    d.save(path)
    return path


def viber(out_vw, cards_dir, batch):
    os.makedirs(out_vw, exist_ok=True)
    lines = [f"# Пакет для Viber і WhatsApp — ГП-{batch.replace('-', '…')}", "",
             "Кожен пост = картинка з файлу + текст нижче. Ритм — вівторок і четвер.", ""]
    for q in GP:
        code = f"ГП-{q['n']:02d}"
        shutil.copy(os.path.join(cards_dir, code, f"{code}_картка-1.png"), os.path.join(out_vw, f"{code}.png"))
        body = q["post"].replace("в боті 👇", "👉 " + link(q["n"]))
        body = "\n".join(l for l in body.split("\n") if l.strip())
        lines += [f"## {code}", f"Картинка: {code}.png", "", "```", body, "```", ""]
    path = os.path.join(out_vw, "Тексти_Viber_WhatsApp.md")
    open(path, "w", encoding="utf-8").write("\n".join(lines))
    return path


if __name__ == "__main__":
    out, batch, date = sys.argv[1], sys.argv[2], sys.argv[3]
    cards = os.path.join(out, f"1_ГП-{batch.replace('-', '–')} картки і відео")
    print(docx_texts(cards, batch, date))
    print(viber(os.path.join(out, f"2_VIBER_WHATSAPP (ГП-{batch.replace('-', '–')})"), cards, batch))
